import io
import logging
import os
import smtplib
from datetime import date, timedelta
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from flask import current_app

logger = logging.getLogger(__name__)


def _build_workbook(reservas) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agenda 7 dias"

    headers = ["DATA", "CONDOMÍNIO", "SALÃO", "APARTAMENTO", "SOLICITANTE", "STATUS"]
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [14, 30, 24, 16, 30, 16]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22

    for row_idx, r in enumerate(reservas, start=2):
        solicitante = (
            "[Festa do Condomínio]" if r.festa_condominio
            else (r.nome_solicitante or "—")
        )
        apartamento = "—" if r.festa_condominio else (r.apartamento or "—")
        ws.append([
            r.data_festa.strftime("%d/%m/%Y"),
            r.salao.condominio.nome,
            r.salao.nome,
            apartamento,
            solicitante,
            r.status.capitalize(),
        ])
        # zebra striping
        if row_idx % 2 == 0:
            fill = PatternFill(start_color="F0F4FA", end_color="F0F4FA", fill_type="solid")
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = fill

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def enviar_backup_agenda():
    """Gera planilha com reservas dos próximos 7 dias e envia por e-mail."""
    try:
        from app.models import Reserva
        from app import db

        hoje = date.today()
        fim  = hoje + timedelta(days=7)

        reservas = (
            db.session.query(Reserva)
            .filter(
                Reserva.data_festa >= hoje,
                Reserva.data_festa <= fim,
                Reserva.status.in_(["confirmado", "pendente"]),
            )
            .order_by(Reserva.data_festa)
            .all()
        )

        buf = _build_workbook(reservas)

        smtp_server = current_app.config.get("SMTP_SERVER") or os.getenv("SMTP_SERVER", "")
        smtp_port   = int(current_app.config.get("SMTP_PORT")   or os.getenv("SMTP_PORT", 587))
        smtp_user   = current_app.config.get("SMTP_USER")   or os.getenv("SMTP_USER", "")
        smtp_pass   = current_app.config.get("SMTP_PASSWORD") or os.getenv("SMTP_PASSWORD", "")
        mail_to     = current_app.config.get("MAIL_TO")     or os.getenv("MAIL_TO", "")

        if not all([smtp_server, smtp_user, smtp_pass, mail_to]):
            logger.warning("Backup de agenda ignorado: credenciais SMTP não configuradas.")
            return

        filename = f"agenda_{hoje.strftime('%Y%m%d')}.xlsx"

        msg = MIMEMultipart()
        msg["From"]    = smtp_user
        msg["To"]      = mail_to
        msg["Subject"] = f"[CondoReservas] Agenda dos próximos 7 dias — {hoje.strftime('%d/%m/%Y')}"

        body = (
            f"Olá,\n\n"
            f"Segue em anexo a agenda de reservas confirmadas e pendentes "
            f"de {hoje.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')}.\n\n"
            f"Total de reservas: {len(reservas)}\n\n"
            f"— CondoReservas (backup automático)"
        )
        msg.attach(MIMEText(body, "plain", "utf-8"))

        attachment = MIMEApplication(buf.read(), _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        attachment.add_header("Content-Disposition", "attachment", filename=filename)
        msg.attach(attachment)

        with smtplib.SMTP(smtp_server, smtp_port, timeout=30) as server:
            server.ehlo()
            server.starttls()
            server.login(smtp_user, smtp_pass)
            server.sendmail(smtp_user, [mail_to], msg.as_string())

        logger.info(f"Backup de agenda enviado para {mail_to} ({len(reservas)} reservas).")

    except Exception as exc:
        logger.error(f"Falha no backup de agenda: {exc}", exc_info=True)
